"""Index styles: the template for Word indexes made from a PDF, and the indexed copy's styles.

    python tests/test_styles.py

Makes a template (.dotm with a fake macro part, and a .dotx) and checks that the index
written from it keeps the template's page setup and styles, with Indicat's settings
(columns, right-aligned numbers and leader, index font) applied on top; and that the
indexed copy turns off "Automatically update document styles"."""
import copy
import sys
import tempfile
import zipfile
from pathlib import Path

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from indicat import static_index as SI       # noqa: E402
from indicat.pipeline import BuildOptions    # noqa: E402
from indicat.settings import DEFAULTS        # noqa: E402

fails = 0


def check(ok, what):
    global fails
    print(("OK   " if ok else "FAIL ") + what)
    if not ok:
        fails += 1


def make_template(path: Path, macro: bool):
    """A template with A5 pages, Garamond 12 pt, an 'index 1' style with its own indent and a
    dotted right tab, some text, and (for .dotm) a macro part."""
    from docx import Document
    from docx.enum.style import WD_STYLE_TYPE
    from docx.enum.text import WD_TAB_ALIGNMENT, WD_TAB_LEADER
    from docx.shared import Inches, Mm, Pt
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Mm(148), Mm(210)
    sec.left_margin = sec.right_margin = Inches(0.75)
    doc.styles["Normal"].font.name = "Garamond"
    doc.styles["Normal"].font.size = Pt(12)
    i1 = doc.styles.add_style("index 1", WD_STYLE_TYPE.PARAGRAPH)
    i1.paragraph_format.left_indent = Inches(0.3)
    i1.paragraph_format.tab_stops.add_tab_stop(Inches(4), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
    doc.add_paragraph("Template body text that must not appear in the index.")
    if macro:              # like a real Normal.dotm: the headings are only latent
        for name in ("Heading 1", "Heading 2"):
            doc.styles[name].delete()
    tmp = path.with_suffix(".tmp.docx")
    doc.save(tmp)
    ctype = (b"application/vnd.ms-word.template.macroEnabledTemplate.main+xml" if macro else
             b"application/vnd.openxmlformats-officedocument.wordprocessingml.template.main+xml")
    with zipfile.ZipFile(tmp) as zin, zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "[Content_Types].xml":
                data = data.replace(
                    b"application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml", ctype)
                if macro:
                    data = data.replace(b"</Types>", b'<Default Extension="bin" ContentType='
                                        b'"application/vnd.ms-office.vbaProject"/></Types>')
            elif item.filename == "word/_rels/document.xml.rels" and macro:
                data = data.replace(b"</Relationships>", b'<Relationship Id="rIdVba" Type="http://schemas.'
                                    b'microsoft.com/office/2006/relationships/vbaProject" Target="vbaProject.bin"/>'
                                    b'</Relationships>')
            zout.writestr(item, data)
        if macro:
            zout.writestr("word/vbaProject.bin", b"not really macros")
    tmp.unlink()


def sample_model():
    E = SI.Entry
    return [SI.Section("Index of Scripture", [("Old Testament", [
        SI.Entry("Genesis", [], [E("1:1", [SI.Run("12")]), E("2:4", [SI.Run("13")])]),
        SI.Entry("Exodus", [], [E("3:14", [SI.Run("40")])])])])]


def test_template(tmpdir: Path, suffix: str, macro: bool):
    from docx import Document
    from docx.enum.text import WD_TAB_ALIGNMENT, WD_TAB_LEADER
    from docx.shared import Inches, Mm
    tmpl = tmpdir / f"House{suffix}"
    make_template(tmpl, macro)
    out = tmpdir / f"Index from {suffix}.docx"
    opts = BuildOptions(out_dir=tmpdir, columns=2, separator="\t", leader="none")
    msgs = SI.write_docx(out, sample_model(), opts, "Brill", tmpl)
    check(not msgs, f"{suffix}: template used without complaint {msgs}")
    with zipfile.ZipFile(out) as z:
        names = z.namelist()
        ctypes = z.read("[Content_Types].xml")
    check(b"wordprocessingml.document.main+xml" in ctypes and b"template" not in ctypes
          and b"macroEnabled" not in ctypes, f"{suffix}: saved as a plain .docx")
    check(not any("vba" in n.lower() for n in names), f"{suffix}: no macro part carried over")
    doc = Document(out)
    text = "\n".join(p.text for p in doc.paragraphs)
    check("Template body text" not in text and "Genesis" in text and "Exodus" in text,
          f"{suffix}: the template's text left out, the index written")
    sec = doc.sections[-1]
    check(abs(sec.page_width - Mm(148)) < Mm(1) and abs(sec.left_margin - Inches(0.75)) < Inches(0.01),
          f"{suffix}: page size and margins from the template")
    check(doc.styles["Normal"].font.size.pt == 12, f"{suffix}: the template's text size kept")
    check(doc.styles["Normal"].font.name == "Brill" and doc.styles["index 1"].font.name == "Brill",
          f"{suffix}: Indicat's index font applied")
    i1 = doc.styles["index 1"].paragraph_format
    check(abs(i1.left_indent - Inches(0.3)) < Inches(0.01), f"{suffix}: the template's indent kept")
    rights = [t for t in i1.tab_stops if t.alignment == WD_TAB_ALIGNMENT.RIGHT]
    check(len(rights) == 1 and rights[0].leader == WD_TAB_LEADER.SPACES,
          f"{suffix}: one right tab, with Indicat's leader (none), not the template's dots")
    cols = [s._sectPr.find("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}cols")
            for s in doc.sections]
    check(any(c is not None and c.get("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}num") == "2"
              for c in cols), f"{suffix}: Indicat's two columns")
    check(doc.styles["Heading 1"] is not None, f"{suffix}: headings available")


def test_choice(tmpdir: Path):
    import os
    s = copy.deepcopy(DEFAULTS)
    old = os.environ.get("APPDATA")
    os.environ["APPDATA"] = str(tmpdir / "appdata")
    try:
        t, note = SI.template_for(s)
        check(t is None and note and "Normal template wasn't found" in note,
              "no Word: falls back to the plain style, with a note")
        normal = tmpdir / "appdata" / "Microsoft" / "Templates" / "Normal.dotm"
        normal.parent.mkdir(parents=True)
        make_template(normal, macro=True)
        t, note = SI.template_for(s)
        check(t == normal and note is None, "default: Word's Normal template")
        s["pdf_template"], s["pdf_template_path"] = "file", str(tmpdir / "missing.dotx")
        t, note = SI.template_for(s)
        check(t is None and "wasn't found" in note, "a chosen template that's gone: plain style, with a note")
        s["pdf_template"] = "plain"
        check(SI.template_for(s) == (None, None), "plain style when chosen")
    finally:
        if old is None:
            os.environ.pop("APPDATA", None)
        else:
            os.environ["APPDATA"] = old


def test_plain(tmpdir: Path):
    from docx import Document
    out = tmpdir / "plain.docx"
    msgs = SI.write_docx(out, sample_model(), BuildOptions(out_dir=tmpdir, separator=", "), "Times New Roman")
    doc = Document(out)
    check(not msgs and doc.styles["Normal"].font.size.pt == 10 and "Genesis" in
          "\n".join(p.text for p in doc.paragraphs), "plain style as before (10 pt)")


def test_link_styles(tmpdir: Path):
    from docx import Document
    from indicat import docx_io as D
    src = tmpdir / "linked.docx"
    Document().save(src)
    with zipfile.ZipFile(src) as z:
        items = [(i, z.read(i.filename)) for i in z.infolist()]
    with zipfile.ZipFile(src, "w") as z:
        for i, data in items:
            if i.filename == "word/settings.xml":
                data = data.replace(b"<w:zoom", b"<w:linkStyles/><w:zoom", 1)
            z.writestr(i, data)
    pkg = D.DocxPackage(src)
    check(D.keep_own_styles(pkg), "“Automatically update document styles” found")
    out = tmpdir / "linked - Indexed.docx"
    pkg.save(out)
    with zipfile.ZipFile(out) as z:
        check(b"linkStyles" not in z.read("word/settings.xml"), "and turned off in the copy")
    with zipfile.ZipFile(src) as z:
        check(b"linkStyles" in z.read("word/settings.xml"), "the original is untouched")
    pkg2 = D.DocxPackage(out)
    check(not D.keep_own_styles(pkg2) and not pkg2.extra_parts, "nothing to do when it is off")


def main():
    with tempfile.TemporaryDirectory() as td:
        td = Path(td)
        test_template(td, ".dotm", macro=True)
        test_template(td, ".dotx", macro=False)
        test_choice(td)
        test_plain(td)
        test_link_styles(td)
    if fails:
        print(f"{fails} failed")
        sys.exit(1)
    print("All style tests passed.")


if __name__ == "__main__":
    main()
