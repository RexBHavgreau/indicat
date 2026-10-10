"""Proofreading list (.docx) and CSV export."""
from __future__ import annotations

import csv
import datetime as _dt
import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, Inches

from . import __version__
from .parser import INDEXED_STATUSES

ROMAN = {"i": 1, "v": 5, "x": 10, "l": 50, "c": 100, "d": 500, "m": 1000}


def roman_value(s: str) -> int | None:
    s = s.lower()
    if not s or any(ch not in ROMAN for ch in s):
        return None
    total = 0
    for i, ch in enumerate(s):
        v = ROMAN[ch]
        if i + 1 < len(s) and ROMAN[s[i + 1]] > v:
            total -= v
        else:
            total += v
    return total


def page_sort_key(o) -> tuple:
    if not o.page:
        return (2, 0, o.seq)
    rv = roman_value(o.page)
    if rv is not None:
        return (0, rv, o.seq)
    m = re.match(r"\d+", o.page)
    return (1, int(m.group()) if m else 0, o.seq)


def _set_cell_shading(cell, hex_fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tcPr.append(shd)


def _repeat_header(row):
    trPr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader")
    el.set(qn("w:val"), "true")
    trPr.append(el)


def _widths(table, widths):
    table.autofit = False
    for i, wd in enumerate(widths):
        table.columns[i].width = Inches(wd)
    for row in table.rows:
        for i, wd in enumerate(widths):
            row.cells[i].width = Inches(wd)


def _context_runs(par, context: str, size=8.5):
    parts = re.split(r"(⟦|⟧)", context)
    bold = False
    for p in parts:
        if p == "⟦":
            bold = True
            continue
        if p == "⟧":
            bold = False
            continue
        r = par.add_run(p)
        r.bold = bold
        r.font.size = Pt(size)


def write_proof(path: Path, csv_path: Path, src: Path, occs, opts, used_word: bool, messages, res=None):
    indexed = [o for o in occs if o.include and o.ref is not None
               and not (o.in_note and opts.footnote_mode == "skip")]
    multi = bool(res is not None and res.defs)
    only_scripture = not multi

    def iname(o):
        return res.index_name(o.index) if res is not None else "Scripture"
    review = [o for o in occs if o not in indexed]
    indexed.sort(key=page_sort_key)

    doc = Document()
    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.LANDSCAPE
    sec.page_width, sec.page_height = Inches(11), Inches(8.5)
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(sec, side, Inches(0.6))
    st = doc.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(10)

    doc.add_heading("Scripture Index — Proofreading List" if only_scripture else "Index — Proofreading List",
                    level=1)
    p = doc.add_paragraph()
    p.add_run("Manuscript: ").bold = True
    p.add_run(src.name)
    p.add_run("\nBuilt: ").bold = True
    p.add_run(_dt.datetime.now().strftime("%Y-%m-%d %H:%M") + f" with Indicat {__version__}")
    p.add_run("\nNotes: ").bold = True
    from .pipeline import FOOTNOTE_MODES
    p.add_run(FOOTNOTE_MODES.get(opts.footnote_mode, opts.footnote_mode))
    p.add_run("\nPage numbers: ").bold = True
    p.add_run("read from Word's layout of the indexed copy on this computer. If fonts, printer, or "
              "text change, rebuild before final proofing." if used_word else
              "not available (Word automation was not used). Entries are listed in reading order.")
    for m in messages:
        q = doc.add_paragraph(style="List Bullet")
        q.add_run(m).italic = True

    doc.add_paragraph(
        "How to use: work through the indexed pages in order. For each row, find the entry on that "
        "page, confirm it, and tick the box. Rows marked “Book from context” are the ones most worth "
        "checking: the book was not named next to the reference." +
        ("" if only_scripture else " Word and term entries show the heading they are filed under."))

    doc.add_heading(f"Indexed entries ({len(indexed)})" if multi else f"Indexed references ({len(indexed)})",
                    level=2)
    if multi:
        cols = ["✓", "Page", "Index", "Index entry", "Where", "How found", "Context"]
        widths = [0.3, 0.5, 0.8, 1.5, 0.95, 1.25, 4.5]
    else:
        cols = ["✓", "Page", "Index entry", "Where", "How the book was found", "Context"]
        widths = [0.3, 0.55, 1.6, 1.05, 1.45, 4.85]
    t = doc.add_table(rows=1, cols=len(cols))
    t.style = "Table Grid"
    hdr = t.rows[0]
    _repeat_header(hdr)
    for i, c in enumerate(cols):
        cell = hdr.cells[i]
        cell.text = ""
        r = cell.paragraphs[0].add_run(c)
        r.bold = True
        r.font.size = Pt(9)
        _set_cell_shading(cell, "D9D9D9")
    for o in indexed:
        row = t.add_row()
        vals = ["☐", o.page or "—", o.ref.display(), o.where, o.status_label]
        if multi:
            vals.insert(2, iname(o))
        how_i, where_i = len(vals) - 1, len(vals) - 2
        if o.merged_warning:
            vals[how_i] += f" — {o.merged_warning}"
        for i, v in enumerate(vals):
            cell = row.cells[i]
            cell.text = ""
            run = cell.paragraphs[0].add_run(v)
            run.font.size = Pt(9)
            if i == how_i and (o.status in ("inferred", "inferred-weak") or o.edited):
                run.font.highlight_color = WD_COLOR_INDEX.YELLOW
            if i == where_i and o.in_note:
                run.bold = True
        _context_runs(row.cells[len(vals)].paragraphs[0], o.context)
    _widths(t, widths)

    doc.add_heading(f"Not indexed — review ({len(review)})", level=2)
    doc.add_paragraph(
        "These were found but left out of the index: verse-only references, chapter-only "
        "references, possible titles, volume:page citations, references in notes when notes are "
        "skipped" + ("" if only_scripture else ", quotations in Greek, Hebrew or other languages") +
        ", and anything you excluded. Add any that belong in the index by hand, or include "
        "them in the app and rebuild.")
    if multi:
        cols2 = ["Reading order", "Index", "Found", "Where", "Reason", "Context"]
        widths2 = [0.7, 0.8, 1.5, 0.95, 1.6, 4.25]
    else:
        cols2 = ["Reading order", "Found", "Where", "Reason", "Context"]
        widths2 = [0.8, 1.6, 1.05, 1.6, 4.75]
    t2 = doc.add_table(rows=1, cols=len(cols2))
    t2.style = "Table Grid"
    _repeat_header(t2.rows[0])
    for i, c in enumerate(cols2):
        cell = t2.rows[0].cells[i]
        cell.text = ""
        r = cell.paragraphs[0].add_run(c)
        r.bold = True
        r.font.size = Pt(9)
        _set_cell_shading(cell, "D9D9D9")
    for o in review:
        row = t2.add_row()
        reason = o.status_label
        if o.status in INDEXED_STATUSES and not o.include:
            reason = "Excluded by you"
        elif o.in_note and opts.footnote_mode == "skip":
            reason = "In a note (notes skipped)"
        if o.hit.reason and o.status != "quotation":
            reason += f" ({o.hit.reason})"
        vals = [str(o.seq + 1), o.display, o.where, reason]
        if multi:
            vals.insert(1, iname(o))
        for i, v in enumerate(vals):
            row.cells[i].text = ""
            row.cells[i].paragraphs[0].add_run(v).font.size = Pt(9)
        _context_runs(row.cells[len(vals)].paragraphs[0], o.context)
    _widths(t2, widths2)
    doc.save(path)

    with open(csv_path, "w", newline="", encoding="utf-8-sig") as f:
        wr = csv.writer(f)
        wr.writerow(["indexed", "page", "index", "book_or_heading", "entry_or_subentry", "where", "status",
                     "reading_order", "context"])
        for o in indexed + review:
            wr.writerow(["yes" if o in indexed else "no", o.page or "", iname(o),
                         o.ref.book_name if o.ref else "", o.ref.label() if o.ref else o.hit.raw,
                         o.where, o.status_label, o.seq + 1,
                         o.context.replace("⟦", "[").replace("⟧", "]")])
