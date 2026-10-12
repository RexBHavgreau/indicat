"""A finished (static) index written straight from the review list: used by the PDF
route, where the page and note numbers are already known.

Layout follows what Word's own index gives (book or heading at level 1, references
or subentries at level 2), with Chicago-style cross-references:
"favor. *See* grace", "Assumption, 9. *See also* Dormition".
Written as Word (.docx) or OpenDocument (.odt).
"""
from __future__ import annotations

import zipfile
from dataclasses import dataclass, field
from pathlib import Path
from xml.sax.saxutils import escape

from . import terms as T
from .word_com import _n_suffix


@dataclass
class Run:
    text: str
    italic: bool = False
    bold: bool = False
    highlight: bool = False


@dataclass
class Entry:
    text: str
    locs: list = field(default_factory=list)      # [Run]
    subs: list = field(default_factory=list)      # [Entry]
    xrefs: list = field(default_factory=list)     # [(see word, [targets])]
    bold: bool = False


@dataclass
class Section:
    title: str
    groups: list = field(default_factory=list)    # [(group title or "", [Entry])]


def _locators(items, mode: str, sep_pages: str = ", ") -> list[Run]:
    """items: [(pdf page, label, note number or None)] -> runs "12, 14n3, 20nn4–5"."""
    by_page: dict = {}
    order = []
    for page, label, note in items:
        if page not in by_page:
            by_page[page] = [label, False, []]
            order.append(page)
        if note is None:
            by_page[page][1] = True
        else:
            by_page[page][2].append(str(note))
    runs: list[Run] = []
    for k, page in enumerate(sorted(order)):
        label, body, notes = by_page[page]
        if k:
            runs.append(Run(sep_pages))
        if body or not notes:
            runs.append(Run(label))                       # a body reference on the page covers its notes
        elif mode == "auto":
            runs.append(Run(label + _n_suffix(notes)))
        elif mode == "bold":
            runs.append(Run(label, bold=True))
        elif mode == "highlight":
            runs.append(Run(label + _n_suffix(notes), highlight=True))
        else:
            runs.append(Run(label, italic=True))
    return runs


def build_model(res, occs, opts) -> list[Section]:
    mode = opts.footnote_mode
    cat = res.catalog
    use = [o for o in occs if o.include and o.ref is not None and not (o.in_note and mode == "skip")]
    sections: list[Section] = []
    # Scripture
    sc = [o for o in use if o.is_scripture]
    if sc and res.scripture:
        sec = Section(opts.heading or cat.index_title)
        groups = cat.groups_in_order()
        titles = dict(groups)
        buckets: dict = {}
        for o in sc:
            o.ref.cat = cat
            g = cat.index_group(o.ref.book) if groups else ""
            buckets.setdefault(g or "", {}).setdefault(o.ref.book, {}).setdefault(
                (o.ref.sort_key(), o.ref.label()), []).append(_item(o))
        order = [L for L, _ in groups] if groups else [""]
        for g in order:
            books = buckets.get(g)
            if not books:
                continue
            entries = []
            for bid in sorted(books, key=lambda b: cat.book_key(b)):
                e = Entry(cat.name(bid), bold=True)
                for (sk, label), items in sorted(books[bid].items()):
                    e.subs.append(Entry(label, _locators(items, mode)))
                entries.append(e)
            sec.groups.append((titles.get(g, ""), entries))
        sections.append(sec)
    # word and term indexes
    see_words = T.SEE_WORDS.get(res.settings.get("language", "en"), T.SEE_WORDS["en"])
    for d in res.defs:
        rows = [o for o in use if o.index == d.id]
        heads: dict = {}
        for o in rows:
            r = o.ref
            hk = (r.sort or T.sort_code(r.heading), r.heading)
            h = heads.setdefault(hk, {"locs": [], "subs": {}})
            if r.sub:
                h["subs"].setdefault((r.sub_sort or T.sort_code(r.sub), r.sub), []).append(_item(o))
            else:
                h["locs"].append(_item(o))
        xrefs: dict = {}
        if d.kind == "terms":
            for ref, kind, target in T.compile_set(d).crossrefs:
                key = (ref.sort or T.sort_code(ref.heading), ref.heading)
                word = see_words[0] if kind == "see" else see_words[1]
                if ref.sub:
                    sk = (ref.sub_sort or T.sort_code(ref.sub), ref.sub)
                    heads.setdefault(key, {"locs": [], "subs": {}})["subs"].setdefault(sk, [])
                    xrefs.setdefault((key, sk), {}).setdefault(word, []).append(target)
                else:
                    heads.setdefault(key, {"locs": [], "subs": {}})
                    xrefs.setdefault((key, None), {}).setdefault(word, []).append(target)
        if not heads:
            continue
        entries = []
        for hk in sorted(heads):
            h = heads[hk]
            e = Entry(hk[1], _locators(h["locs"], mode))
            e.xrefs = sorted(xrefs.get((hk, None), {}).items(), key=lambda x: len(x[0]))
            for sk in sorted(h["subs"]):
                s = Entry(sk[1], _locators(h["subs"][sk], mode))
                s.xrefs = sorted(xrefs.get((hk, sk), {}).items(), key=lambda x: len(x[0]))
                e.subs.append(s)
            entries.append(e)
        sections.append(Section(d.title, [("", entries)]))
    return sections


def _item(o):
    return (getattr(o, "pdf_page", 0) or 0, o.page or "", o.note_number if o.in_note else None)


def entry_runs(e: Entry, sep: str) -> list[Run]:
    """The text of one index line: entry, separator, page numbers, cross-references (CMOS)."""
    runs = [Run(e.text)]
    if e.locs:
        runs.append(Run(sep))
        runs.extend(e.locs)
    for word, targets in e.xrefs:
        runs.append(Run(". "))
        runs.append(Run(word, italic=True))
        runs.append(Run(" " + "; ".join(sorted(dict.fromkeys(targets), key=T.fold_str))))
    return runs


# ---------------------------------------------------------------------------
# Word
# ---------------------------------------------------------------------------
def write_docx(path: Path, sections: list[Section], opts, font: str) -> None:
    from docx import Document
    from docx.enum.section import WD_SECTION
    from docx.enum.text import WD_COLOR_INDEX, WD_TAB_ALIGNMENT, WD_TAB_LEADER
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt

    doc = Document()
    st = doc.styles
    normal = st["Normal"]
    normal.font.name = font
    normal.font.size = Pt(10)
    rpr = normal.element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts")
        rpr.insert(0, rf)
    for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rf.set(qn(a), font)
    normal.paragraph_format.space_after = Pt(0)
    for hn in ("Heading 1", "Heading 2"):
        st[hn].font.name = font
        st[hn].font.color.rgb = None
    sec0 = doc.sections[0]
    text_w = sec0.page_width - sec0.left_margin - sec0.right_margin
    cols = max(1, int(opts.columns or 1))
    col_w = int((text_w - (cols - 1) * Inches(0.5)) / cols)
    tab = opts.separator == "\t"
    leader = {"dot": WD_TAB_LEADER.DOTS, "none": WD_TAB_LEADER.SPACES}.get(opts.leader or "none")

    def style(name, left, bold=False):
        try:
            s = st[name]
        except KeyError:
            from docx.enum.style import WD_STYLE_TYPE
            s = st.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
            s.base_style = normal
        pf = s.paragraph_format
        pf.left_indent = Inches(left)
        pf.first_line_indent = Inches(-0.2)
        pf.space_after = Pt(0)
        s.font.bold = bold
        if tab:
            pf.tab_stops.add_tab_stop(col_w - Inches(0.01), WD_TAB_ALIGNMENT.RIGHT, leader)
        return s

    s1 = style("index 1", 0.2)
    s2 = style("index 2", 0.4)

    def set_cols(section, n):
        sp = section._sectPr
        c = sp.find(qn("w:cols"))
        if c is None:
            c = OxmlElement("w:cols")
            sp.append(c)
        c.set(qn("w:num"), str(n))
        c.set(qn("w:space"), "720")

    def para(text_runs, style_obj, bold=False):
        p = doc.add_paragraph(style=style_obj)
        for r in text_runs:
            run = p.add_run(r.text)
            if r.italic:
                run.italic = True
            if r.bold or bold:
                run.bold = True
            if r.highlight:
                run.font.highlight_color = WD_COLOR_INDEX.YELLOW
        return p

    first = True
    for sec in sections:
        if not first:
            ns = doc.add_section(WD_SECTION.NEW_PAGE if opts.new_page else WD_SECTION.CONTINUOUS)
            set_cols(ns, 1)
        first = False
        doc.add_heading(sec.title, level=1)
        body = doc.add_section(WD_SECTION.CONTINUOUS)
        set_cols(body, cols)
        for gtitle, entries in sec.groups:
            if gtitle:
                h = doc.add_heading(gtitle, level=2)
                h.paragraph_format.keep_with_next = True
            for e in entries:
                p = para(entry_runs(e, opts.separator), s1, bold=e.bold)
                if e.subs and len(e.locs) <= 12:
                    p.paragraph_format.keep_with_next = True
                for s in e.subs:
                    para(entry_runs(s, opts.separator), s2)
    doc.save(path)


# ---------------------------------------------------------------------------
# OpenDocument
# ---------------------------------------------------------------------------
_ODT_NS = ('xmlns:office="urn:oasis:names:tc:opendocument:xmlns:office:1.0" '
           'xmlns:style="urn:oasis:names:tc:opendocument:xmlns:style:1.0" '
           'xmlns:text="urn:oasis:names:tc:opendocument:xmlns:text:1.0" '
           'xmlns:fo="urn:oasis:names:tc:opendocument:xmlns:xsl-fo-compatible:1.0" '
           'xmlns:svg="urn:oasis:names:tc:opendocument:xmlns:svg-compatible:1.0" '
           'xmlns:loext="urn:org:documentfoundation:names:experimental:office:xmlns:loext:1.0" '
           'office:version="1.3"')


def _odt_text(s: str) -> str:
    """Escape, keeping tabs and runs of spaces as ODF elements."""
    out = []
    for part in s.split("\t"):
        out.append(escape(part).replace("  ", " <text:s/>"))
    return "<text:tab/>".join(out)


def write_odt(path: Path, sections: list[Section], opts, font: str) -> None:
    cols = max(1, int(opts.columns or 1))
    text_w = 6.5
    col_w = (text_w - (cols - 1) * 0.5) / cols
    tab = opts.separator == "\t"
    leader = {"dot": ".", "none": " "}.get(opts.leader or "none", " ")
    tabs = (f'<style:tab-stops><style:tab-stop style:position="{col_w - 0.25:.3f}in" style:type="right" '
            f'style:leader-text="{leader}"/></style:tab-stops>') if tab else ""
    fq = escape(font, {'"': "&quot;"})
    styles = f'''<?xml version="1.0" encoding="UTF-8"?>
<office:document-styles {_ODT_NS}>
<office:font-face-decls><style:font-face style:name="{fq}" svg:font-family="&apos;{fq}&apos;"/></office:font-face-decls>
<office:styles>
<style:default-style style:family="paragraph"><style:paragraph-properties fo:margin-top="0in" fo:margin-bottom="0in"/>
<style:text-properties style:font-name="{fq}" fo:font-size="10pt"/></style:default-style>
<style:style style:name="Standard" style:family="paragraph" style:class="text"/>
<style:style style:name="Index_20_Title" style:display-name="Index Title" style:family="paragraph" style:parent-style-name="Standard">
<style:paragraph-properties fo:margin-bottom="0.17in" fo:keep-with-next="always"/>
<style:text-properties fo:font-size="16pt" fo:font-weight="bold"/></style:style>
<style:style style:name="Index_20_Group" style:display-name="Index Group" style:family="paragraph" style:parent-style-name="Standard">
<style:paragraph-properties fo:margin-top="0.12in" fo:margin-bottom="0.06in" fo:keep-with-next="always"/>
<style:text-properties fo:font-size="12pt" fo:font-weight="bold"/></style:style>
<style:style style:name="Index_20_1" style:display-name="Index 1" style:family="paragraph" style:parent-style-name="Standard">
<style:paragraph-properties fo:margin-left="0.2in" fo:text-indent="-0.2in">{tabs}</style:paragraph-properties></style:style>
<style:style style:name="Index_20_2" style:display-name="Index 2" style:family="paragraph" style:parent-style-name="Standard">
<style:paragraph-properties fo:margin-left="0.4in" fo:text-indent="-0.2in">{tabs}</style:paragraph-properties></style:style>
</office:styles>
<office:automatic-styles><style:page-layout style:name="pm1"><style:page-layout-properties fo:page-width="8.5in" fo:page-height="11in" fo:margin-top="1in" fo:margin-bottom="1in" fo:margin-left="1in" fo:margin-right="1in"/></style:page-layout></office:automatic-styles>
<office:master-styles><style:master-page style:name="Standard" style:page-layout-name="pm1"/></office:master-styles>
</office:document-styles>'''
    auto = [f'<style:style style:name="Sect1" style:family="section"><style:section-properties text:dont-balance-text-columns="true">'
            f'<style:columns fo:column-count="{cols}" fo:column-gap="0.5in"/></style:section-properties></style:style>',
            '<style:style style:name="Ti" style:family="text"><style:text-properties fo:font-style="italic"/></style:style>',
            '<style:style style:name="Tb" style:family="text"><style:text-properties fo:font-weight="bold"/></style:style>',
            '<style:style style:name="Th" style:family="text"><style:text-properties fo:background-color="#ffff00"/></style:style>',
            '<style:style style:name="PBreak" style:family="paragraph" style:parent-style-name="Index_20_Title">'
            '<style:paragraph-properties fo:break-before="page"/></style:style>',
            '<style:style style:name="P1b" style:family="paragraph" style:parent-style-name="Index_20_1">'
            '<style:paragraph-properties fo:keep-with-next="always"/><style:text-properties fo:font-weight="bold"/></style:style>',
            '<style:style style:name="P1k" style:family="paragraph" style:parent-style-name="Index_20_1">'
            '<style:paragraph-properties fo:keep-with-next="always"/></style:style>']

    def runs_xml(runs):
        out = []
        for r in runs:
            t = _odt_text(r.text)
            sty = "Ti" if r.italic else "Tb" if r.bold else "Th" if r.highlight else None
            out.append(f'<text:span text:style-name="{sty}">{t}</text:span>' if sty else t)
        return "".join(out)

    body = []
    for k, sec in enumerate(sections):
        tsty = "PBreak" if (k and opts.new_page) else "Index_20_Title"
        body.append(f'<text:p text:style-name="{tsty}">{escape(sec.title)}</text:p>')
        body.append(f'<text:section text:style-name="Sect1" text:name="Index{k + 1}">')
        for gtitle, entries in sec.groups:
            if gtitle:
                body.append(f'<text:p text:style-name="Index_20_Group">{escape(gtitle)}</text:p>')
            for e in entries:
                # keep a heading with its first subentry only when the heading line is short (a long
                # paragraph kept with the next one loses its text in LibreOffice's column sections)
                sty = "P1b" if e.bold else ("P1k" if e.subs and len(e.locs) <= 12 else "Index_20_1")
                body.append(f'<text:p text:style-name="{sty}">{runs_xml(entry_runs(e, opts.separator))}</text:p>')
                for s in e.subs:
                    body.append(f'<text:p text:style-name="Index_20_2">{runs_xml(entry_runs(s, opts.separator))}</text:p>')
        body.append('</text:section>')
    content = (f'<?xml version="1.0" encoding="UTF-8"?>\n<office:document-content {_ODT_NS}>'
               f'<office:font-face-decls><style:font-face style:name="{fq}" svg:font-family="&apos;{fq}&apos;"/>'
               f'</office:font-face-decls><office:automatic-styles>{"".join(auto)}</office:automatic-styles>'
               f'<office:body><office:text>{"".join(body)}</office:text></office:body></office:document-content>')
    manifest = ('<?xml version="1.0" encoding="UTF-8"?>\n'
                '<manifest:manifest xmlns:manifest="urn:oasis:names:tc:opendocument:xmlns:manifest:1.0" '
                'manifest:version="1.3">'
                '<manifest:file-entry manifest:full-path="/" manifest:media-type="application/vnd.oasis.opendocument.text"/>'
                '<manifest:file-entry manifest:full-path="content.xml" manifest:media-type="text/xml"/>'
                '<manifest:file-entry manifest:full-path="styles.xml" manifest:media-type="text/xml"/>'
                '<manifest:file-entry manifest:full-path="meta.xml" manifest:media-type="text/xml"/>'
                '</manifest:manifest>')
    meta = ('<?xml version="1.0" encoding="UTF-8"?>\n<office:document-meta '
            'xmlns:office="urn:oasis:names:tc:opendocument:xmlns:office:1.0" '
            'xmlns:meta="urn:oasis:names:tc:opendocument:xmlns:meta:1.0" office:version="1.3">'
            '<office:meta><meta:generator>Indicat</meta:generator></office:meta></office:document-meta>')
    with zipfile.ZipFile(path, "w") as z:
        z.writestr(zipfile.ZipInfo("mimetype"), "application/vnd.oasis.opendocument.text",
                   compress_type=zipfile.ZIP_STORED)
        z.writestr("META-INF/manifest.xml", manifest, compress_type=zipfile.ZIP_DEFLATED)
        z.writestr("styles.xml", styles, compress_type=zipfile.ZIP_DEFLATED)
        z.writestr("content.xml", content, compress_type=zipfile.ZIP_DEFLATED)
        z.writestr("meta.xml", meta, compress_type=zipfile.ZIP_DEFLATED)
